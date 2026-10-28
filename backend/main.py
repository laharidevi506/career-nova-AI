import io
import os
import json
import re
import time
import traceback
import pymupdf
from docx import Document
from dotenv import load_dotenv
from google import genai

from fastapi import (
    FastAPI,
    HTTPException,
    UploadFile,
    File,
    Form,
    Depends,
)
from fastapi.middleware.cors import CORSMiddleware

from sqlalchemy.orm import Session
from pwdlib import PasswordHash
from pydantic import BaseModel, Field

import models
from database import engine, Base, SessionLocal
from schemas import UserRegister, UserLogin
# =========================================================
# LOAD ENVIRONMENT VARIABLES
# =========================================================

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

# Change the model here if your Gemini API account
# does not support this model.
GEMINI_MODEL = "gemini-3.1-flash-lite"

gemini_client = (
    genai.Client(api_key=GEMINI_API_KEY)
    if GEMINI_API_KEY
    else None
)


# =========================================================
# CREATE DATABASE TABLES
# =========================================================

Base.metadata.create_all(bind=engine)


# =========================================================
# INITIALIZE FASTAPI
# =========================================================

app = FastAPI(
    title="Career Nova AI",
    description="AI-Powered Career Mentor and Placement Readiness Platform",
    version="1.0.0",
)


# =========================================================
# CORS CONFIGURATION
# =========================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# PASSWORD HASHING
# =========================================================

password_hash = PasswordHash.recommended()


# =========================================================
# DATABASE SESSION
# =========================================================

def get_db():
    db = SessionLocal()

    try:
        yield db
    finally:
        db.close()


# =========================================================
# CAREER GUIDANCE REQUEST SCHEMA
# =========================================================

class CareerGuidanceRequest(BaseModel):
    education: str = Field(min_length=1, max_length=100)
    interests: str = Field(min_length=1, max_length=1000)
    skills: str = Field(min_length=1, max_length=1000)
    goals: str = Field(min_length=1, max_length=1000)
# NEW SCHEMA: Skill Gap Analysis
class SkillGapRequest(BaseModel):
    current_skills: str = Field(min_length=1, max_length=2000)
    target_career: str = Field(min_length=2, max_length=200)
    education: str = Field(min_length=2, max_length=200)
    experience_level: str = Field(min_length=1, max_length=50)
   #personalized roadmap request
class RoadmapRequest(BaseModel):
    target_career: str = Field(min_length=2, max_length=200)
    current_skills: str = Field(min_length=1, max_length=2000)
    daily_study_hours: int = Field(ge=1, le=12)
    duration_weeks: int = Field(ge=4, le=52)
    education: str = Field(min_length=2, max_length=200)
###Placement preparation
class PlacementPreparationRequest(BaseModel):
    target_role: str = Field(min_length=2, max_length=200)
    education: str = Field(min_length=2, max_length=200)
    experience_level: str = Field(min_length=2, max_length=50)
    programming_language: str = Field(min_length=2, max_length=50)
    interview_type: str = Field(min_length=2, max_length=50)

class ChatbotRequest(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    history: list[dict] = Field(default_factory=list)   

# =========================================================
# HOME ENDPOINT
# =========================================================

@app.get("/")
def home():
    return {
        "message": "Welcome to Career Nova AI Backend!",
        "status": "Running successfully",
    }


# =========================================================
# HEALTH ENDPOINT
# =========================================================

@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "ai_configured": gemini_client is not None,
    }


# =========================================================
# USER REGISTRATION ENDPOINT
# =========================================================

@app.post("/register", status_code=201)
def register_user(
    user: UserRegister,
    db: Session = Depends(get_db),
):
    existing_user = (
        db.query(models.User)
        .filter(models.User.email == user.email)
        .first()
    )

    if existing_user:
        raise HTTPException(
            status_code=409,
            detail="Email is already registered",
        )

    hashed_password = password_hash.hash(user.password)

    new_user = models.User(
        full_name=user.full_name,
        email=user.email,
        hashed_password=hashed_password,
        role="student",
        education=user.education,
        interests=user.interests,
        skills=user.skills,
        goals=user.goals,
    )

    try:
        db.add(new_user)
        db.commit()
        db.refresh(new_user)

    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Registration failed. Please try again.",
        )

    return {
        "message": "User registered successfully!",
        "user": {
            "id": new_user.id,
            "full_name": new_user.full_name,
            "email": new_user.email,
            "role": new_user.role,
            "education": new_user.education,
            "interests": new_user.interests,
            "skills": new_user.skills,
            "goals": new_user.goals,
        },
    }


# =========================================================
# USER LOGIN ENDPOINT
# =========================================================

@app.post("/login")
def login_user(
    user: UserLogin,
    db: Session = Depends(get_db),
):
    existing_user = (
        db.query(models.User)
        .filter(models.User.email == user.email)
        .first()
    )

    if not existing_user:
        raise HTTPException(
            status_code=401,
            detail="Invalid email or password",
        )

    if not password_hash.verify(
        user.password,
        existing_user.hashed_password,
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid email or password",
        )

    return {
        "message": "Login successful!",
        "user": {
            "id": existing_user.id,
            "full_name": existing_user.full_name,
            "email": existing_user.email,
            "role": existing_user.role,
            "education": existing_user.education,
            "interests": existing_user.interests,
            "skills": existing_user.skills,
            "goals": existing_user.goals,
        },
    }

# =========================================================
# MODULE 1: AI CAREER GUIDANCE
# =========================================================

@app.post("/career-guidance")
def career_guidance(request: CareerGuidanceRequest):

    # Check required fields
    education = request.education.strip()
    interests = request.interests.strip()
    skills = request.skills.strip()
    goals = request.goals.strip()

    if not all([education, interests, skills, goals]):
        raise HTTPException(
            status_code=422,
            detail="Please complete all career guidance fields.",
        )

    # Check whether Gemini is configured
    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail=(
                "Gemini API key is missing. "
                "Check the GEMINI_API_KEY in backend/.env."
            ),
        )

    # Build the AI prompt
    prompt = f"""
You are Career Nova AI, a helpful career guidance assistant
for college students and fresh graduates.

Analyze the following student's profile:

Education: {education}
Interests: {interests}
Current skills: {skills}
Career goals: {goals}

Recommend exactly three suitable career paths.

For each career, provide:
1. title: career name
2. description: a clear description of the role
3. reason: explain why it matches this student's profile
4. required_skills: a list of five relevant skills
5. project: one practical beginner-friendly project

Requirements:
- Personalize recommendations using the student's profile.
- Do not claim the student has skills they did not mention.
- Be realistic and encouraging.
- Do not guarantee jobs, salaries, or placements.
- Use simple, student-friendly English.
- Return valid JSON only.
- Do not include Markdown fences or text outside the JSON.

Use this exact JSON structure:

{{
  "careers": [
    {{
      "title": "Career name",
      "description": "Career description",
      "reason": "Personalized explanation",
      "required_skills": [
        "Skill 1",
        "Skill 2",
        "Skill 3",
        "Skill 4",
        "Skill 5"
      ],
      "project": "Project description"
    }}
  ]
}}
"""

    try:
        # Call Gemini and retry temporary 503/overload errors.
        response = None

        for attempt in range(3):
            try:
                response = gemini_client.models.generate_content(
                    model=GEMINI_MODEL,
                    contents=prompt,
                    config={
                        "response_mime_type": "application/json",
                        "temperature": 0.4,
                    },
                )
                break

            except Exception as exc:
                error_message = str(exc).lower()
                is_temporary_error = (
                    "503" in error_message
                    or "unavailable" in error_message
                    or "overloaded" in error_message
                    or "high demand" in error_message
                )

                if is_temporary_error and attempt < 2:
                    wait_seconds = 2 ** (attempt + 1)
                    print(
                        f"Gemini is temporarily busy. "
                        f"Retrying in {wait_seconds} seconds "
                        f"(attempt {attempt + 1}/2)...",
                        flush=True,
                    )
                    time.sleep(wait_seconds)
                else:
                    raise

        if response is None or not response.text:
            raise HTTPException(
                status_code=502,
                detail="Gemini returned an empty response. Please try again.",
            )

        # Convert AI response to Python data
        ai_data = json.loads(response.text)

        careers = ai_data.get("careers", [])

        # Validate the response structure
        if not isinstance(careers, list) or len(careers) != 3:
            raise ValueError("Gemini returned an unexpected career list.")

        for career in careers:
            if not isinstance(career, dict) or not all(
                key in career
                for key in [
                    "title",
                    "description",
                    "reason",
                    "required_skills",
                    "project",
                ]
            ):
                raise ValueError("Gemini returned incomplete career details.")

            if not isinstance(career["required_skills"], list):
                raise ValueError("Invalid required_skills format.")

        return {
            "message": "AI career recommendations generated successfully!",
            "careers": careers,
        }

    except HTTPException:
        raise

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=502,
            detail="Gemini returned an invalid response. Please try again.",
        )

    except ValueError:
        raise HTTPException(
            status_code=502,
            detail="Gemini returned an unexpected response format.",
        )

    except Exception as exc:
        import traceback

        print("GEMINI ERROR DETAILS:", repr(exc), flush=True)
        traceback.print_exc()

        raise HTTPException(
            status_code=502,
            detail=(
                "AI career guidance is temporarily unavailable. "
                "Please try again shortly and check the backend terminal "
                "if the problem continues."
            ),
        )

# =========================================================
# MODULE 2: SKILL GAP ANALYSIS
# =========================================================

@app.post("/skill-gap-analysis")
def skill_gap_analysis(request: SkillGapRequest):
    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="Gemini API key is missing. Check backend/.env.",
        )

    current_skills = request.current_skills.strip()
    target_career = request.target_career.strip()
    education = request.education.strip()
    experience_level = request.experience_level.strip()

    if not all([current_skills, target_career, education, experience_level]):
        raise HTTPException(
            status_code=422,
            detail="Please complete all Skill Gap Analysis fields.",
        )

    prompt = f"""
You are Career Nova AI, an expert career counsellor for college students.
Compare the student's reported skills with the skills normally needed for
this target career. Be realistic and personalize every recommendation.

Student information:
Education: {education}
Current skills: {current_skills}
Target career: {target_career}
Experience level: {experience_level}

Return one JSON object with exactly this structure:
{{
  "summary": "A personalized summary of the student's readiness",
  "current_strengths": ["A strength supported by the student's information"],
  "missing_skills": [
    {{
      "skill": "Skill name",
      "importance": "Why the skill matters",
      "priority": "High",
      "reason": "Why the student should learn it",
      "learning_resource": "A useful learning resource, preferably with its official name"
    }}
  ],
  "action_plan": [
    {{"task": "A practical learning task", "duration": "2 weeks"}}
  ]
}}

Requirements:
- Include 3 to 6 missing skills, 2 to 5 current strengths, and 4 to 6 action steps.
- Each priority must be exactly High, Medium, or Low.
- Do not claim the student has skills they did not report.
- Suggest realistic learning resources and practical tasks.
- Do not guarantee employment, salary, or placement.
- Use simple, student-friendly English.
- Return valid JSON only. Do not include Markdown fences or commentary.
"""

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={
                "response_mime_type": "application/json",
                "temperature": 0.2,
            },
        )

        raw_text = getattr(response, "text", None)
        if not raw_text or not raw_text.strip():
            print("SKILL GAP: Gemini returned no text.", flush=True)
            raise HTTPException(
                status_code=502,
                detail="Gemini returned an empty response. Please try again.",
            )

        # Log the response for debugging. Never log the API key.
        print("SKILL GAP GEMINI RAW RESPONSE:", repr(raw_text), flush=True)

        # Parse JSON directly first. If the model wraps it in Markdown fences
        # or adds a short preamble, clean/extract the outer JSON object.
        cleaned_text = raw_text.strip()
        if cleaned_text.startswith("```"):
            cleaned_text = re.sub(
                r"^```(?:json)?\s*", "", cleaned_text, flags=re.IGNORECASE
            )
            cleaned_text = re.sub(r"\s*```$", "", cleaned_text).strip()

        try:
            result = json.loads(cleaned_text)
        except json.JSONDecodeError:
            first_brace = cleaned_text.find("{")
            last_brace = cleaned_text.rfind("}")
            if first_brace == -1 or last_brace <= first_brace:
                raise
            result = json.loads(cleaned_text[first_brace:last_brace + 1])

        if not isinstance(result, dict):
            raise ValueError("Gemini response must be a JSON object")

        required_keys = ["summary", "current_strengths", "missing_skills", "action_plan"]
        for key in required_keys:
            if key not in result:
                raise ValueError(f"Missing required field: {key}")

        if not isinstance(result["summary"], str):
            raise ValueError("Invalid summary format")
        if not isinstance(result["current_strengths"], list):
            raise ValueError("Invalid current_strengths format")
        if not isinstance(result["missing_skills"], list):
            raise ValueError("Invalid missing_skills format")
        if not isinstance(result["action_plan"], list):
            raise ValueError("Invalid action_plan format")

        for skill in result["missing_skills"]:
            if not isinstance(skill, dict) or not all(
                key in skill
                for key in ["skill", "importance", "priority", "reason", "learning_resource"]
            ):
                raise ValueError("Invalid missing_skills item")
            if skill["priority"] not in ["High", "Medium", "Low"]:
                skill["priority"] = "Medium"

        for step in result["action_plan"]:
            if not isinstance(step, dict) or not all(
                key in step for key in ["task", "duration"]
            ):
                raise ValueError("Invalid action_plan item")

        return result

    except HTTPException:
        raise
    except json.JSONDecodeError as exc:
        print("SKILL GAP JSON PARSE ERROR:", str(exc), flush=True)
        raise HTTPException(
            status_code=502,
            detail=(
                "Gemini did not return valid JSON. Please retry once. "
                "If the problem continues, check the backend terminal."
            ),
        )
    except ValueError as exc:
        print("SKILL GAP RESPONSE VALIDATION ERROR:", str(exc), flush=True)
        raise HTTPException(
            status_code=502,
            detail="Gemini returned an unexpected response format. Please try again.",
        )
    except Exception as exc:
        import traceback
        print("SKILL GAP ANALYSIS ERROR:", repr(exc), flush=True)
        traceback.print_exc()
        raise HTTPException(
            status_code=502,
            detail="Skill Gap Analysis is temporarily unavailable. Please try again shortly.",
        )
    
   ###Personalize road mapp---------------------------------------

@app.post("/personalized-roadmap")
def personalized_roadmap(data: RoadmapRequest):
    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="Gemini AI is not configured."
        )

    prompt = f"""
    You are an expert career mentor.

    Create a personalized learning roadmap using:
    Target career: {data.target_career}
    Current skills: {data.current_skills}
    Education: {data.education}
    Daily study hours: {data.daily_study_hours}
    Duration: {data.duration_weeks} weeks

    Return ONLY valid JSON with this structure:
    {{
        "career": "Target career",
        "overview": "Personalized roadmap summary",
        "total_duration": "{data.duration_weeks} weeks",
        "daily_study_hours": {data.daily_study_hours},
        "phases": [
            {{
                "phase": 1,
                "title": "Phase title",
                "week_range": "Weeks 1-2",
                "goals": ["Learning goal"],
                "topics": ["Topic"],
                "tasks": ["Practical task"],
                "project": "Project description",
                "resources": ["Free learning resource"],
                "milestone": "Expected achievement"
            }}
        ],
        "final_projects": [
            {{
                "title": "Project title",
                "description": "Project description",
                "skills_used": ["Python"]
            }}
        ],
        "career_tips": ["Career tip"]
    }}

    Generate enough phases to cover the entire duration.
    Make the plan realistic for the student's daily study hours.
    """

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={
                "response_mime_type": "application/json",
                "temperature": 0.4
            }
        )

        if not response.text:
            raise HTTPException(
                status_code=502,
                detail="Gemini returned an empty response."
            )

        result = json.loads(response.text)

        return result

    except json.JSONDecodeError:
        raise HTTPException(
            status_code=502,
            detail="Gemini returned invalid JSON. Please try again."
        )

    except HTTPException:
        raise

    except Exception as e:
        print("Personalized Roadmap Error:", repr(e))
        raise HTTPException(
            status_code=502,
            detail="Unable to generate roadmap. Please try again."
        )


@app.post("/placement-preparation")
def placement_preparation(request: PlacementPreparationRequest):
    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="AI service is not configured."
        )

    prompt = f"""
You are an expert technical interviewer and placement preparation coach.

Prepare a personalized interview preparation guide for this student.

Student details:
- Target role: {request.target_role}
- Education: {request.education}
- Experience level: {request.experience_level}
- Programming language: {request.programming_language}
- Interview type: {request.interview_type}

Return ONLY valid JSON using this exact structure:

{{
    "target_role": "{request.target_role}",
    "overview": "A short personalized preparation overview",
    "technical_questions": [
        {{
            "question": "Technical interview question",
            "sample_answer": "Simple, clear sample answer",
            "explanation": "Why this answer is correct",
            "difficulty": "Beginner"
        }}
    ],
    "coding_questions": [
        {{
            "question": "Coding problem",
            "approach": "Step-by-step approach",
            "sample_solution": "Working Python code with explanation",
            "difficulty": "Beginner"
        }}
    ],
    "hr_questions": [
        {{
            "question": "HR interview question",
            "sample_answer": "Sample answer the student can personalize",
            "tips": ["Useful interview tip"]
        }}
    ],
    "preparation_tips": [
        "Practical placement preparation tip"
    ]
}}

Requirements:
1. Provide 5 technical questions.
2. Provide 3 coding questions appropriate for the student's level.
3. Provide 5 HR interview questions.
4. Give clear, accurate, beginner-friendly answers.
5. Use the selected programming language for coding solutions where appropriate.
6. Tailor the questions to the target role and interview type.
7. Include useful tips without guaranteeing placement success.
8. Return valid JSON only, without Markdown fences.
"""

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={
                "response_mime_type": "application/json",
                "temperature": 0.4
            }
        )

        result = json.loads(response.text)

        return result

    except json.JSONDecodeError as exc:
        print("Placement Preparation JSON error:", exc)
        raise HTTPException(
            status_code=502,
            detail="AI returned an invalid response. Please try again."
        )

    except Exception as exc:
        print("Placement Preparation error:", repr(exc))
        raise HTTPException(
            status_code=502,
            detail="Unable to generate placement preparation. Please try again."
        )

###resume analyzer

@app.post("/resume-analyzer")
async def analyze_resume(
    file: UploadFile = File(...),
    target_role: str = Form(...)
):
    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="AI service is not configured."
        )

    filename = (file.filename or "").lower()

    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Upload a PDF or DOCX resume."
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    if len(file_bytes) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail="Resume must be smaller than 5 MB."
        )

    try:
        if filename.endswith(".pdf"):
            with pymupdf.open(
                stream=file_bytes,
                filetype="pdf"
            ) as pdf:
                resume_text = "\n".join(
                    page.get_text() for page in pdf
                )
        else:
            document = Document(io.BytesIO(file_bytes))
            resume_text = "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
            )

            for table in document.tables:
                for row in table.rows:
                    resume_text += "\n" + " | ".join(
                        cell.text for cell in row.cells
                    )

    except Exception as exc:
        print("Resume extraction error:", repr(exc))
        raise HTTPException(
            status_code=400,
            detail="Unable to read this file. Please upload a valid PDF or DOCX."
        )

    if len(resume_text.strip()) < 50:
        raise HTTPException(
            status_code=400,
            detail="Not enough readable text found. Try a text-based PDF or DOCX resume."
        )

    prompt = f"""
You are an experienced resume reviewer and career coach.

Analyze the resume for this target role: {target_role}

RESUME TEXT:
{resume_text[:18000]}

Return valid JSON only, with this exact structure:

{{
  "target_role": "{target_role}",
  "resume_score": 75,
  "summary": "Short summary of the resume and its suitability for the target role.",
  "strengths": [
    "Specific strength supported by the resume"
  ],
  "weaknesses": [
    "Specific weakness or missing information"
  ],
  "missing_skills": [
    {{
      "skill": "Relevant skill",
      "importance": "High",
      "reason": "Why this skill matters for the target role"
    }}
  ],
  "ats_feedback": [
    "Practical ATS compatibility observation"
  ],
  "improvement_suggestions": [
    "Specific and actionable improvement"
  ],
  "projects_feedback": [
    "Feedback about projects and how to improve their descriptions"
  ]
}}

Rules:
- resume_score must be an integer between 0 and 100.
- Base feedback on the supplied resume and target role.
- Do not invent qualifications, experience, or achievements.
- Distinguish missing information from confirmed lack of a skill.
- Give constructive, actionable feedback.
- Evaluate ATS compatibility using the available text only; do not claim to have tested an actual ATS.
- Return JSON only, without Markdown fences.
"""

    try:
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={
                "response_mime_type": "application/json",
                "temperature": 0.3
            }
        )

        result = json.loads(response.text)

        score = result.get("resume_score")
        if not isinstance(score, int) or not 0 <= score <= 100:
            raise ValueError("Invalid resume score returned by AI")

        return result

    except json.JSONDecodeError as exc:
        print("Resume Analyzer JSON error:", repr(exc))
        raise HTTPException(
            status_code=502,
            detail="AI returned an invalid response. Please try again."
        )

    except Exception as exc:
        print("Resume Analyzer error:", repr(exc))
        raise HTTPException(
            status_code=502,
            detail="Unable to analyze the resume. Please try again."
        )

# =========================================================
# MODULE 6: PROGRESS TRACKING
# =========================================================

from datetime import datetime


# Request schema for adding a learning task
class ProgressTaskRequest(BaseModel):
    user_id: int = Field(gt=0)
    task_title: str = Field(min_length=1, max_length=200)
    category: str = Field(default="General", min_length=1, max_length=100)


# Helper function to format task data
def format_progress_task(task):
    return {
        "id": task.id,
        "user_id": task.user_id,
        "task_title": task.task_title,
        "category": task.category,
        "is_completed": task.is_completed,
        "study_date": (
            task.study_date.isoformat()
            if task.study_date else None
        ),
        "created_at": (
            task.created_at.isoformat()
            if task.created_at else None
        ),
        "completed_at": (
            task.completed_at.isoformat()
            if task.completed_at else None
        ),
    }


# GET: Display all learning tasks for a user
@app.get("/progress/{user_id}")
def get_progress(
    user_id: int,
    db: Session = Depends(get_db),
):
    user = db.query(models.User).filter(
        models.User.id == user_id
    ).first()

    if not user:
        raise HTTPException(
            status_code=404,
            detail="User not found."
        )

    tasks = (
        db.query(models.LearningProgress)
        .filter(models.LearningProgress.user_id == user_id)
        .order_by(models.LearningProgress.id.desc())
        .all()
    )

    return {
        "tasks": [
            format_progress_task(task)
            for task in tasks
        ]
    }


# POST: Add a new learning task
@app.post("/progress", status_code=201)
def add_progress_task(
    request: ProgressTaskRequest,
    db: Session = Depends(get_db),
):
    user = db.query(models.User).filter(
        models.User.id == request.user_id
    ).first()

    if not user:
        raise HTTPException(
            status_code=404,
            detail="User not found."
        )

    task_title = request.task_title.strip()
    category = request.category.strip()

    if not task_title or not category:
        raise HTTPException(
            status_code=422,
            detail="Task name and category cannot be empty."
        )

    task = models.LearningProgress(
        user_id=request.user_id,
        task_title=task_title,
        category=category,
        is_completed=False,
    )

    try:
        db.add(task)
        db.commit()
        db.refresh(task)

        return {
            "message": "Learning task added successfully!",
            "task": format_progress_task(task),
        }

    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Unable to add task. Please try again."
        )


# PATCH: Mark a task completed or incomplete
@app.patch("/progress/{task_id}/toggle")
def toggle_progress_task(
    task_id: int,
    db: Session = Depends(get_db),
):
    task = db.query(models.LearningProgress).filter(
        models.LearningProgress.id == task_id
    ).first()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Learning task not found."
        )

    task.is_completed = not task.is_completed

    if task.is_completed:
        task.completed_at = datetime.now()
    else:
        task.completed_at = None

    try:
        db.commit()
        db.refresh(task)

        return {
            "message": "Progress updated successfully!",
            "task": format_progress_task(task),
        }

    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Unable to update progress."
        )


# DELETE: Delete a learning task
@app.delete("/progress/{task_id}")
def delete_progress_task(
    task_id: int,
    db: Session = Depends(get_db),
):
    task = db.query(models.LearningProgress).filter(
        models.LearningProgress.id == task_id
    ).first()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Learning task not found."
        )

    try:
        db.delete(task)
        db.commit()

        return {
            "message": "Learning task deleted successfully!"
        }

    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Unable to delete task."
        )
        # =========================================================
# MODULE 7: INTERNSHIP PORTAL
# =========================================================

@app.get("/internships")
def get_internships(
    search: str = "",
    location: str = "",
    db: Session = Depends(get_db),
):
    try:
        # Show only verified and open internships.
        query = db.query(models.Internship).filter(
            models.Internship.opportunity_status == "open",
            models.Internship.verification_status == "verified",
        )

        # Search by internship title or company name.
        if search.strip():
            search_term = f"%{search.strip()}%"

            query = query.filter(
                (models.Internship.title.ilike(search_term))
                | (
                    models.Internship.company_name.ilike(
                        search_term
                    )
                )
            )

        # Filter internships by location.
        if location.strip():
            query = query.filter(
                models.Internship.location.ilike(
                    f"%{location.strip()}%"
                )
            )

        internships = (
            query.order_by(models.Internship.id.desc()).all()
        )

        # Format the results for the frontend.
        internship_list = [
            {
                "id": item.id,
                "title": item.title,
                "company_name": item.company_name,
                "description": item.description,
                "required_skills": item.required_skills,
                "location": item.location,
                "work_mode": item.work_mode,
                "stipend": item.stipend,
                "application_url": item.application_url,
                "source_url": item.source_url,
                "deadline": (
                    item.deadline.isoformat()
                    if item.deadline
                    else None
                ),
            }
            for item in internships
        ]

        return {
            "total": len(internship_list),
            "internships": internship_list,
        }

    except Exception as e:
        import traceback

        print(
            "INTERNSHIP PORTAL ERROR:",
            repr(e),
            flush=True,
        )
        traceback.print_exc()

        raise HTTPException(
            status_code=500,
            detail="Unable to load internships. Check the backend terminal.",
        )
    

# =========================================================
# MODULE 8: AI CAREER CHATBOT - FAST VERSION
# =========================================================

@app.post("/chatbot")
def career_chatbot(request: ChatbotRequest):

    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="AI service is not configured. Check backend/.env."
        )

    message = request.message.strip()

    if not message:
        raise HTTPException(
            status_code=422,
            detail="Please enter a message."
        )

    # Keep only the latest 4 conversation messages
    conversation = []

    for item in request.history[-4:]:
        if not isinstance(item, dict):
            continue

        role = item.get("role")
        content = item.get("content")

        if (
            role in ("user", "assistant")
            and isinstance(content, str)
            and content.strip()
        ):
            conversation.append(
                f"{role.capitalize()}: {content[:1000]}"
            )

    prompt = f"""
You are Career Nova AI, a helpful career and learning assistant.

Help students with:
- Career guidance
- Python and programming
- Coding problems and algorithms
- Resume improvement
- Technical and HR interviews
- Skills and learning roadmaps

Use simple English. Give direct answers with examples when useful.
For simple questions, answer briefly.
Do not claim to access live job listings without verified live data.

Recent conversation:
{chr(10).join(conversation)}

Student question:
{message}

Answer the student's question directly.
"""

    try:
        # One attempt only to avoid long retry delays
        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config={
                "temperature": 0.3,
                "max_output_tokens": 700,
            },
        )

        answer = getattr(response, "text", None)

        if not answer or not answer.strip():
            raise HTTPException(
                status_code=502,
                detail="AI returned an empty answer. Please try again."
            )

        return {"reply": answer.strip()}

    except HTTPException:
        raise

    except Exception as exc:
        print("CHATBOT ERROR:", repr(exc), flush=True)

        error_message = str(exc).lower()

        if any(word in error_message for word in (
            "503", "unavailable", "overloaded",
            "high demand", "429", "resource_exhausted",
            "rate limit"
        )):
            raise HTTPException(
                status_code=503,
                detail=(
                    "Gemini AI is busy right now. "
                    "Please wait a few seconds and try again."
                ),
            ) from exc

        raise HTTPException(
            status_code=502,
            detail="Unable to generate an answer. Please try again."
        ) from exc

